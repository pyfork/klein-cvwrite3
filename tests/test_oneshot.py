"""Chunk 3: profile and upload, one ad in -> one CV out, the retry, checkpoints, the pause, edit and re-render.

Claude is the fake (app/fakeclaude.py). Renders use the real Chrome and poppler, so these tests skip without them.
"""
import io
import shutil
import unittest
import zipfile
from pathlib import Path

import psycopg

from app import adread, cvtext, ic, pause, upload, worker
from render import build_cv
from tests.support import AD, CVCase, AppCase, OWNER

HAVE_RENDER = Path(build_cv.CHROME).exists() and shutil.which("pdfinfo") and shutil.which("pdftotext")
needs_render = unittest.skipUnless(HAVE_RENDER, "needs Chrome and poppler")
FAKE_IC = "920315" + "-10-" + "5241"  # made up, date-valid


class Units(unittest.TestCase):
    def test_ic_numbers(self):
        self.assertEqual(ic.find(f"name\nIC: {FAKE_IC}\n"), [(2, "92••••-••-••41")])
        self.assertEqual(len(ic.find("9203151" + "052411")), 0)  # 13 digits: not an IC
        self.assertEqual(len(ic.find("921315" + "105241")), 0)  # month 13
        self.assertEqual(len(ic.find("920315 10 5241")), 1)

    def test_quote_check(self):
        ad = "We need someone who’s calm — really calm.\nKeywords: SQL, route optimi-\nsation, stakeholder reporting"
        read = adread.verify({
            "job_title": "Analyst", "company": "",
            "problems": [{"problem": "calm", "quote": "who's calm - really calm"}, {"problem": "x", "quote": "not in the ad"}],
            "keywords_top": ["SQL", "route optimisation", "Power BI"], "keywords_secondary": ["stakeholder reporting"],
            "red_flags": [{"flag": "invented", "quote": "unlimited overtime"}],
        }, ad)
        self.assertEqual([p["problem"] for p in read["problems"]], ["calm"])
        self.assertEqual(read["keywords_top"], ["SQL", "route optimisation", "stakeholder reporting"])
        self.assertEqual(read["red_flags"], [])
        with self.assertRaises(adread.AdReadError):
            adread.verify({"problems": [{"problem": "x", "quote": "nope"}], "keywords_top": ["SQL"]}, ad)

    def test_word_check(self):
        user = "SUMMARY\nI fix late deliveries with SQL and routes.\n\nEXPERIENCE\n- Built a report.\n" * 5
        ok, kept = cvtext.word_check(user, "## SUMMARY\n" + user.replace("SUMMARY\n", "", 1))
        self.assertTrue(ok)
        reworded = user.replace("fix late deliveries", "resolve tardy shipments")
        self.assertFalse(cvtext.word_check(user, "## SUMMARY\n" + reworded)[0])
        padded = user + "\nA whole new paragraph the model added on its own initiative here.\n"
        self.assertFalse(cvtext.word_check(user, "## SUMMARY\n" + padded)[0])

    def test_plain_round_trip(self):
        md = "# {{NAME}}\n**Analyst**\nA tagline\n{{CONTACT}}\n\n---\n\n## PROFESSIONAL SUMMARY\nShort **bold** text.\n\n## SKILLS\n- **SQL:** daily\n"
        plain = cvtext.to_plain(md)
        self.assertTrue(plain.startswith("[NAME]\n[CONTACT]"))
        self.assertNotIn("**", plain)
        self.assertNotIn("#", plain)
        body, changed = cvtext.plain_body(plain)
        self.assertTrue(body.startswith("PROFESSIONAL SUMMARY"))
        self.assertFalse(changed)
        self.assertTrue(cvtext.plain_body("Jane Doe\nnew phone\n" + body)[1])
        self.assertEqual(cvtext.summary_words(md), 3)


class Upload(unittest.TestCase):
    TEXT = ("Test Person\nOperations analyst.\n" + "Built a daily SQL report for hub managers. " * 10)

    def test_text_files(self):
        self.assertIn("SQL report", upload.read("cv.md", self.TEXT.encode()))
        self.assertIn("SQL report", upload.read("cv.TXT", self.TEXT.encode("cp1252")))

    def test_docx_with_a_table(self):
        import docx
        d = docx.Document()
        d.add_paragraph(self.TEXT)
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Skills", "Route optimisation"
        buf = io.BytesIO()
        d.save(buf)
        text = upload.read("cv.docx", buf.getvalue())
        self.assertIn("Skills | Route optimisation", text)

    def test_refusals(self):
        for name, data, msg in (
            ("cv.doc", b"x" * 300, "Old .doc"), ("cv.exe", b"x" * 300, "Use a .docx"),
            ("cv.md", b"x" * (upload.MAX_BYTES + 1), "over 2 MB"), ("cv.md", b"short", "almost no text"),
            ("cv.docx", b"not a zip at all" * 30, "damaged"),
        ):
            with self.assertRaises(upload.UploadError) as e:
                upload.read(name, data)
            self.assertIn(msg, str(e.exception))

    def test_zip_bomb(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"0" * (upload.MAX_UNZIPPED + 1))
        self.assertLess(len(buf.getvalue()), upload.MAX_BYTES)
        with self.assertRaises(upload.UploadError):
            upload.read("cv.docx", buf.getvalue())


class Profile(AppCase):
    EMAIL = "nadia@example.com"

    def setUp(self):
        super().setUp()
        self.approve(self.EMAIL)
        self.agree(self.EMAIL)

    def test_upload_previews_without_saving(self):
        from app import auth
        files = {"file": ("cv.md", ("Test Person\n" + "Built a daily SQL report. " * 20).encode(), "text/markdown")}
        r = self.client.post("/profile/upload", data={"csrf": auth.csrf_token(self.settings, self.EMAIL)}, files=files,
                             headers={**self.as_user(self.EMAIL), "Origin": self.settings.app_origin})
        self.assertEqual(r.status_code, 200)
        self.assertIn("Nothing is saved yet", r.text)
        with self.conn() as c:
            self.assertIsNone(c.execute("select 1 from profiles").fetchone())

    def test_ic_number_blocks_save(self):
        r = self.post("/profile", self.EMAIL, {"master_text": f"Name\nIC {FAKE_IC}\n" + "text " * 60, "name": "N"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("Line 2: 92••••-••-••41", r.text)
        self.assertNotIn(FAKE_IC, r.text.replace(f"IC {FAKE_IC}", ""))  # only the textarea echoes it back
        with self.conn() as c:
            self.assertIsNone(c.execute("select 1 from profiles").fetchone())

    def test_generate_needs_a_profile(self):
        r = self.get("/generate", self.EMAIL)
        self.assertIn("Upload your master CV first", r.text)
        self.assertEqual(self.post("/generate", self.EMAIL, {"ad": AD}).status_code, 409)


@needs_render
class Generate(CVCase):
    def test_one_ad_in_one_cv_out(self):
        run_id = self.start()
        self.assertIn("Working on it", self.get(f"/runs/{run_id}", self.EMAIL).text)
        self.drain()
        run = self.run_row(run_id)
        self.assertEqual(run["status"], "passed")
        self.assertEqual(run["drafts"], 1)
        self.assertEqual(self.credits(), 1)
        with self.conn() as c:
            spend = c.execute("select * from credit_ledger where kind = 'spend'").fetchone()
            self.assertEqual((spend["delta"], spend["run_id"]), (-1, run_id))
            steps = [r["step"] for r in c.execute("select step from run_steps where run_id = %s order by id", (run_id,))]
            self.assertEqual(steps, ["queue_wait", "call_a", "call_b", "render", "check"])
            usage = c.execute("select step, model, cost_usd from usage order by id").fetchall()
            self.assertEqual([u["step"] for u in usage], ["call_a", "call_b"])
            self.assertTrue(all(u["cost_usd"] > 0 for u in usage))
        page = self.get(f"/runs/{run_id}", self.EMAIL).text
        self.assertIn("Passed", page)
        self.assertIn("our on-time rate isn&#39;t keeping up", page)  # quoted from the ad
        self.assertNotIn("this sentence is not in the ad", page)  # the fake's invented quote was dropped
        pdf = self.get(f"/runs/{run_id}/cv.pdf", self.EMAIL)
        self.assertEqual((pdf.status_code, pdf.content[:5]), (200, b"%PDF-"))
        md = self.get(f"/runs/{run_id}/cv.md", self.EMAIL)
        for r in (pdf, md, self.get(f"/runs/{run_id}", self.EMAIL)):  # never cached at Cloudflare's edge
            self.assertEqual(r.headers.get("cache-control"), "private, no-store")
        self.assertNotIn("no-store", self.client.get("/static/calm.css").headers.get("cache-control", ""))
        md = md.text
        self.assertIn("# Nadia Hamzah", md)
        self.assertNotIn("{{", md)

    def test_claude_never_sees_contact_values(self):
        seen = []
        orig = self.claude.write_cv
        self.claude.write_cv = lambda msg: (seen.append(msg), orig(msg))[1]
        self.start()
        self.drain()
        self.assertTrue(seen)
        for value in ("Nadia Hamzah", "nadia@example.com", "12-345 6789"):
            self.assertNotIn(value, seen[0])

    def test_retry_then_pass_uses_one_credit(self):
        self.claude.script = ["write:3pages"]
        run_id = self.start()
        self.drain()
        run = self.run_row(run_id)
        self.assertEqual((run["status"], run["drafts"]), ("passed", 2))
        self.assertEqual(self.credits(), 1)

    def test_both_drafts_fail_no_credit_used(self):
        self.claude.script = ["write:3pages", "write:3pages"]
        run_id = self.start()
        self.drain()
        run = self.run_row(run_id)
        self.assertEqual((run["status"], run["failure"]), ("failed", "check_failed"))
        self.assertEqual(self.credits(), 2)
        self.assertIn("passing CV for this ad. No credit used.", self.get(f"/runs/{run_id}", self.EMAIL).text)
        self.assertEqual(self.get(f"/runs/{run_id}/cv.pdf", self.EMAIL).status_code, 404)

    def test_long_summary_is_caught_before_chrome(self):
        self.claude.script = ["write:longsummary"]
        run_id = self.start()
        self.drain()
        with self.conn() as c:
            steps = [r["step"] for r in c.execute("select step from run_steps where run_id = %s order by id", (run_id,))]
        self.assertEqual(steps, ["queue_wait", "call_a", "call_b", "retry", "render", "check"])
        self.assertEqual(self.run_row(run_id)["status"], "passed")

    def test_checkpoints_mean_no_second_claude_call(self):
        from app import pipeline
        orig = pipeline._render_and_check
        crashes = []

        def crash_once(*a, **kw):  # the worker dies after call B, before the render
            if not crashes:
                crashes.append(1)
                raise RuntimeError("worker died")
            return orig(*a, **kw)
        pipeline._render_and_check = crash_once
        try:
            run_id = self.start()
            self.drain()
        finally:
            pipeline._render_and_check = orig
        self.assertEqual(self.claude.calls, ["read", "write"])  # once each, then resumed from the checkpoint
        self.assertEqual(self.run_row(run_id)["status"], "passed")
        self.assertEqual(self.credits(), 1)

    def test_a_run_is_charged_once(self):
        run_id = self.start()
        self.drain()
        with self.conn() as c:
            acct = c.execute("select account_id from runs").fetchone()["account_id"]
            with self.assertRaises(psycopg.errors.UniqueViolation):
                c.execute("insert into credit_ledger (account_id, delta, kind, run_id) values (%s, -1, 'spend', %s)", (acct, run_id))

    def test_missing_keyword_is_a_gap_not_a_failure(self):
        self.claude.script = ["write:missingkw"]
        run_id = self.start()
        self.drain()
        run = self.run_row(run_id)
        self.assertEqual((run["status"], run["drafts"]), ("passed", 1))  # no retry spent forcing it
        missing = [k for k, v in run["keyword_status"].items() if v == "missing"]
        self.assertEqual(len(missing), 1)
        self.assertNotIn(missing[0], run["checks"]["keywords"])
        page = self.get(f"/runs/{run_id}", self.EMAIL).text
        self.assertIn("isn&#39;t in your CV", page)
        self.assertIn("Edit the text</a> (free, no credit)", page)

    def test_unreported_keywords_stay_required(self):
        from app import pipeline
        self.assertEqual(pipeline.keyword_status(["SQL", "WMS"], [{"keyword": "sql", "status": "missing"}, {"keyword": "x", "status": "odd"}]),
                         {"SQL": "missing", "WMS": "backed"})

    def test_claude_error_fails_with_no_credit(self):
        self.claude.script = ["read:error"]
        run_id = self.start()
        self.drain()
        self.assertEqual(self.run_row(run_id)["status"], "failed")
        self.assertEqual(self.credits(), 2)

    def test_others_cant_see_a_run(self):
        run_id = self.start()
        self.approve("mallory@example.com")
        self.agree("mallory@example.com")
        self.assertEqual(self.get(f"/runs/{run_id}", "mallory@example.com").status_code, 404)

    def test_balance_is_reserved_while_generating(self):
        self.start()
        self.start()
        self.assertEqual(self.post("/generate", self.EMAIL, {"ad": AD}).status_code, 409)

    def test_cancelled_run_is_never_charged(self):
        run_id = self.start()
        with self.conn() as c:
            c.execute("update runs set status = 'cancelled' where id = %s", (run_id,))
        self.drain()
        self.assertEqual(self.credits(), 2)


@needs_render
class Pause(CVCase):
    def test_paused_refuses_new_and_holds_saved(self):
        early = self.start()
        with self.conn() as c:
            pause.pause(c, "test", "owner")
        self.assertEqual(self.post("/generate", self.EMAIL, {"ad": AD}).status_code, 409)
        self.assertIn("Generating is paused", self.get("/generate", self.EMAIL).text)
        self.drain()
        self.assertEqual(self.run_row(early)["status"], "failed")  # not yet at call A: refused, no credit
        with self.conn() as c:
            pause.resume(c)
        late = self.start()
        with self.conn() as c:  # as if call B had finished just before the pause
            job = worker.claim(c)
        worker.run_job(self.settings, self.claude, job)  # passes (not paused yet)
        self.assertEqual(self.run_row(late)["status"], "passed")

    def test_job_past_call_b_is_held_then_finished(self):
        run_id = self.start()
        orig = self.claude.write_cv

        def write_then_pause(msg):
            r = orig(msg)
            with self.conn() as c:
                pause.pause(c, "test", "owner")
            return r
        self.claude.write_cv = write_then_pause
        self.drain()
        with self.conn() as c:
            self.assertEqual(c.execute("select status from jobs where run_id = %s", (run_id,)).fetchone()["status"], "held")
            pause.resume(c)
        self.claude.write_cv = orig
        self.claude.calls.clear()
        self.drain()
        self.assertEqual(self.run_row(run_id)["status"], "passed")
        self.assertEqual(self.claude.calls, [])

    def test_claude_errors_from_two_users_pause(self):
        self.approve("bob@example.com")
        self.agree("bob@example.com")
        self.post("/profile", "bob@example.com", {"master_text": (Path(__file__).parent / "fixtures/master_cv.md").read_text(), "name": "Bob"})
        self.claude.script = ["read:error"] * 5
        for i in range(5):
            self.start(email=self.EMAIL if i % 2 else "bob@example.com")
            self.drain()
        with self.conn() as c:
            self.assertEqual(pause.current(c)["reason"], "Claude API errors")

    def test_owner_controls(self):
        self.post("/admin/pause", OWNER)
        self.assertIn("Resume", self.get("/admin", OWNER).text)
        self.post("/admin/test-render", OWNER)
        self.drain()
        self.assertIn("2 pages", self.get("/admin", OWNER).text)
        self.post("/admin/resume", OWNER)
        with self.conn() as c:
            self.assertIsNone(pause.current(c))


@needs_render
class Rerender(CVCase):
    def setUp(self):
        super().setUp()
        self.run_id = self.start()
        self.drain()

    def edit(self, change=lambda t: t):
        text = self.get(f"/runs/{self.run_id}/edit", self.EMAIL)
        self.assertEqual(text.status_code, 200)
        box = cvtext.to_plain(self.run_row(self.run_id)["draft_md"])
        r = self.post(f"/runs/{self.run_id}/edit", self.EMAIL, {"text": change(box)})
        self.drain()
        return r

    def test_edit_rerenders_and_counts(self):
        self.edit(lambda t: t.replace("Eleven years", "Twelve years"))
        run = self.run_row(self.run_id)
        self.assertEqual(run["rerenders_used"], 1)
        self.assertIn("Twelve years", run["draft_md"])
        self.assertEqual(run["pdf_rev"], 2)
        self.assertIn("# Nadia Hamzah", self.get(f"/runs/{self.run_id}/cv.md", self.EMAIL).text)

    def test_header_changes_are_ignored(self):
        self.edit(lambda t: t.replace("[NAME]", "Somebody Else"))
        run = self.run_row(self.run_id)
        self.assertIn("{{NAME}}", run["draft_md"])
        self.assertNotIn("Somebody Else", run["draft_md"])

    def test_rewrite_fails_the_word_check_and_doesnt_count(self):
        self.claude.script = ["reformat:rewrite"]
        self.edit()
        self.assertEqual(self.run_row(self.run_id)["rerenders_used"], 0)
        page = self.get(f"/runs/{self.run_id}/edit", self.EMAIL).text
        self.assertIn("3 of 3 free re-renders left", page)

    def test_three_pages_blocks_and_counts(self):
        self.edit(lambda t: t.replace("KEY ACHIEVEMENTS", "KEY ACHIEVEMENTS\n" + "\n".join(f"- Extra line {i} to fill a third page with many words." for i in range(70))))
        run = self.run_row(self.run_id)
        self.assertEqual(run["rerenders_used"], 1)
        self.assertEqual(run["pdf_rev"], 1)
        with self.conn() as c:
            self.assertIn("runs to 3 pages", c.execute("select message from jobs where kind = 'rerender'").fetchone()["message"])

    def test_limit_of_three(self):
        with self.conn() as c:
            c.execute("update runs set rerenders_used = 3 where id = %s", (self.run_id,))
        r = self.post(f"/runs/{self.run_id}/edit", self.EMAIL, {"text": "SUMMARY\nx"})
        self.assertEqual(r.status_code, 409)
        self.assertIn("Start a new CV for 1 credit", r.text)

    def test_ic_number_in_an_edit_is_stopped_before_claude(self):
        self.claude.calls.clear()
        r = self.post(f"/runs/{self.run_id}/edit", self.EMAIL, {"text": f"SUMMARY\nIC {FAKE_IC}"})
        self.assertEqual(r.status_code, 400)
        self.drain()
        self.assertEqual(self.claude.calls, [])

    def test_missing_keyword_only_warns(self):
        self.edit(lambda t: t.replace("Strong SQL", "Strong databases"))
        with self.conn() as c:
            job = c.execute("select outcome, message from jobs where kind = 'rerender'").fetchone()
        self.assertEqual(job["outcome"], "passed")
        self.assertIn("No longer in your CV", job["message"])


class Prompts(unittest.TestCase):
    def test_loader_reads_the_stub_folder(self):
        from app import prompts
        d = str(Path(__file__).parent / "fixtures" / "stub_prompts")
        for name in prompts.NAMES:
            self.assertIn("Stub", prompts.get(name, d))
        with self.assertRaises(KeyError):
            prompts.get("other", d)


if __name__ == "__main__":
    unittest.main()
